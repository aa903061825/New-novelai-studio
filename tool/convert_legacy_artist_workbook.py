from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


HEADER_VALUES = {"名称", "画风对照"}
FEATURE_LABELS = {"特征", "角色特征", "羊角，耳朵"}
UPPER_PROMPT_LABELS = {"上装", "上半身"}
LOWER_PROMPT_LABELS = {"下装", "下半身"}
OUTFIT_LABELS = {
    "服装",
    "衣服",
}


def clean(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\xa0", " ").strip()


def is_sensitive(value: str) -> bool:
    normalized = value.strip().lower()
    return (
        normalized == "api"
        or normalized.startswith("bearer ")
        or "pst-" in normalized
        or "api_key" in normalized
    )


def looks_like_prompt(value: str) -> bool:
    if not value or is_sensitive(value):
        return False
    lowered = value.lower()
    return "artist:" in lowered or "::" in value or "_" in value or "," in value


def looks_like_character_prompt(value: str) -> bool:
    if not value or is_sensitive(value):
        return False
    lowered = value.lower()
    if "artist:" in lowered:
        return False
    return (
        "::" in value
        or "_" in value
        or "," in value
        or bool(re.search(r"\([^)\n]+\)", value))
    )


NEGATIVE_PROMPT_TERMS = {
    "bad quality",
    "low quality",
    "worst quality",
    "bad pupils",
    "bad hands",
    "bad eyes",
    "bad anatomy",
    "missing finger",
    "extra finger",
    "extra digits",
    "fewer digits",
    "bad face",
    "duplicate",
    "mutation",
    "deformed",
    "disfigured",
    "extra arms",
    "extra legs",
    "bad feet",
    "bad proportions",
    "censored",
    "mosaic",
    "blurry",
    "lowres",
    "artistic error",
    "jpeg artifacts",
    "logo",
    "watermark",
    "signature",
    "text",
}

NON_ARTIST_STYLE_TERMS = {
    "official art",
    "official_style",
    "natural",
    "humid skin",
    "detailed skin texture",
    "shiny skin",
    "deep skin",
    "motion lines",
    "speed lines",
    "great quality",
    "sd",
    "chibi",
    "jaggy lines",
    "outline",
    "detailed",
}


def prompt_parts(value: str) -> list[str]:
    return [
        part.strip()
        for part in re.split(r"[,\n]", value.replace("\r\n", "\n"))
        if part.strip()
    ]


def normalized_prompt_text(value: str) -> str:
    return re.sub(r"[_\s]+", " ", value.lower())


def negative_prompt_score(value: str) -> int:
    normalized = normalized_prompt_text(value)
    return sum(1 for term in NEGATIVE_PROMPT_TERMS if term in normalized)


def clean_prompt_part(value: str) -> str:
    part = value.strip()
    part = re.sub(r"^[{\[\s]+", "", part)
    part = re.sub(r"[}\]\s]+$", "", part)
    part = re.sub(r"^\d+(?:\.\d+)?::\s*", "", part)
    part = re.sub(r"\s*::$", "", part)
    return part.strip(" :")


def is_non_artist_style_part(value: str) -> bool:
    cleaned = normalized_prompt_text(clean_prompt_part(value))
    if not cleaned:
        return True
    if cleaned in NON_ARTIST_STYLE_TERMS:
        return True
    if cleaned.startswith("year "):
        return True
    return negative_prompt_score(cleaned) > 0


def unprefixed_artist_name_count(value: str) -> int:
    count = 0
    for part in prompt_parts(value):
        cleaned = clean_prompt_part(part)
        lowered = cleaned.lower()
        if "artist:" in lowered or is_non_artist_style_part(cleaned):
            continue
        if re.search(r"[_()]|^[a-z][a-z0-9!]*(?:[ _-][a-z0-9!]+){0,3}$", cleaned, re.I):
            count += 1
    return count


def artist_token_count(value: str) -> int:
    return len(re.findall(r"\bartist\s*:", value, flags=re.I))


def artist_prompt_like(value: str) -> bool:
    if not value or is_sensitive(value):
        return False
    artist_count = artist_token_count(value)
    negative_score = negative_prompt_score(value)
    if negative_score >= 8 and negative_score >= max(artist_count, 1) * 4:
        return False
    if artist_count > 0:
        return True
    return unprefixed_artist_name_count(value) >= 3


def stable_id(prefix: str, sheet: str, row: int, column: int = 0) -> str:
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", sheet).strip("-").lower()
    if not slug:
        slug = "sheet"
    if column:
        return f"legacy-{prefix}-{slug}-r{row}c{column}"
    return f"legacy-{prefix}-{slug}-r{row}"


def source_ref(sheet: str, row: int, column: int | None = None) -> dict[str, Any]:
    return {
        "type": "legacy_artist_workbook",
        "id": f"{sheet}!R{row}" + (f"C{column}" if column else ""),
    }


def guess_artist_name(rows: list[list[str]], row: int, column: int) -> str:
    candidates = (
        (row, column + 1),
        (row, column - 1),
        (row - 1, column + 1),
        (row - 1, column),
        (row + 1, column + 1),
        (row + 1, column),
    )
    for candidate_row, candidate_column in candidates:
        if not (0 <= candidate_row < len(rows)):
            continue
        if not (0 <= candidate_column < len(rows[candidate_row])):
            continue
        value = rows[candidate_row][candidate_column]
        if (
            value
            and value not in HEADER_VALUES
            and not looks_like_prompt(value)
            and not is_sensitive(value)
        ):
            return value[:80]
    return ""


def artist_name_fallback(index: int) -> str:
    return f"画师串 {index + 1:02d}"


def import_artist_chains(workbook) -> list[dict[str, Any]]:
    if "画风" not in workbook.sheetnames:
        return []

    worksheet = workbook["画风"]
    rows = [[clean(cell.value) for cell in row] for row in worksheet.iter_rows()]
    artist_chains: list[dict[str, Any]] = []
    seen_prompts: set[str] = set()
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            if (
                not artist_prompt_like(value)
                or value in HEADER_VALUES
                or value in seen_prompts
            ):
                continue
            seen_prompts.add(value)
            name = guess_artist_name(rows, row_index, column_index)
            if not name:
                name = artist_name_fallback(len(artist_chains))
            artist_chains.append(
                {
                    "id": stable_id(
                        "artist",
                        worksheet.title,
                        row_index + 1,
                        column_index + 1,
                    ),
                    "name": name,
                    "tags": [name] if name else [],
                    "prompt": value,
                    "source": source_ref(
                        worksheet.title,
                        row_index + 1,
                        column_index + 1,
                    ),
                }
            )
    return artist_chains


def is_label(value: str) -> bool:
    return (
        value in FEATURE_LABELS
        or value in UPPER_PROMPT_LABELS
        or value in LOWER_PROMPT_LABELS
        or value in OUTFIT_LABELS
    )


def promptish_pair_values(row: list[str], start: int) -> list[tuple[str, str, int]]:
    pairs: list[tuple[str, str, int]] = []
    index = start
    while index < len(row):
        label = row[index]
        value = row[index + 1] if index + 1 < len(row) else ""
        if label and value and is_label(label):
            pairs.append((label, value, index + 1))
            index += 2
            continue
        if label and looks_like_character_prompt(label):
            pairs.append((f"变体 {len(pairs) + 1}", label, index))
        index += 1
    return pairs


def outfit_from_prompt(id_: str, name: str, prompt: str) -> dict[str, Any]:
    return {
        "id": id_,
        "name": name,
        "tags": {"front_full": prompt},
    }


def import_header_character(
    worksheet,
    labels: list[str],
    values: list[str],
) -> dict[str, Any] | None:
    if len(values) < 2 or not values[1]:
        return None

    name = labels[1] or "角色 1"
    feature_parts = [values[1]]
    upper_parts: list[str] = []
    lower_parts: list[str] = []
    outfits: list[dict[str, Any]] = []
    for index, (label, value) in enumerate(zip(labels, values)):
        if not value or is_sensitive(value):
            continue
        if label in FEATURE_LABELS:
            feature_parts.append(value)
            continue
        if label in UPPER_PROMPT_LABELS:
            upper_parts.append(value)
            continue
        if label in LOWER_PROMPT_LABELS:
            lower_parts.append(value)
            continue
        if index == 1:
            continue
        if not label:
            continue
        if looks_like_character_prompt(value):
            outfits.append(
                outfit_from_prompt(
                    f"{stable_id('character', worksheet.title, 1)}-outfit-{index}",
                    label or f"变体 {len(outfits) + 1}",
                    value,
                )
            )

    return {
        "id": stable_id("character", worksheet.title, 1),
        "name": name,
        "traits": {
            "front_full": ", ".join(feature_parts),
            **({"front_upper": ", ".join(upper_parts)} if upper_parts else {}),
            **({"front_lower": ", ".join(lower_parts)} if lower_parts else {}),
        },
        "outfits": outfits,
        "defaultOutfitId": outfits[0]["id"] if outfits else None,
        "source": source_ref(worksheet.title, 1),
    }


def import_row_character(worksheet, row_number: int, row: list[str]) -> dict[str, Any] | None:
    name = row[0] if row else ""
    if not name or is_sensitive(name) or looks_like_prompt(name):
        return None
    if len(row) < 2 or not row[1] or is_sensitive(row[1]):
        return None
    if not looks_like_character_prompt(row[1]):
        return None

    feature_parts = [row[1]]
    upper_parts: list[str] = []
    lower_parts: list[str] = []
    outfits: list[dict[str, Any]] = []
    pair_values = promptish_pair_values(row, 2)
    for label, value, source_column in pair_values:
        if is_sensitive(value):
            continue
        if label in FEATURE_LABELS:
            feature_parts.append(value)
            continue
        if label in UPPER_PROMPT_LABELS:
            upper_parts.append(value)
            continue
        if label in LOWER_PROMPT_LABELS:
            lower_parts.append(value)
            continue
        outfit_id = f"{stable_id('character', worksheet.title, row_number)}-outfit-{source_column + 1}"
        outfits.append(outfit_from_prompt(outfit_id, label, value))

    result = {
        "id": stable_id("character", worksheet.title, row_number),
        "name": name,
        "traits": {
            "front_full": ", ".join(feature_parts),
            **({"front_upper": ", ".join(upper_parts)} if upper_parts else {}),
            **({"front_lower": ", ".join(lower_parts)} if lower_parts else {}),
        },
        "outfits": outfits,
        "source": source_ref(worksheet.title, row_number),
    }
    if outfits:
        result["defaultOutfitId"] = outfits[0]["id"]
    return result


def import_characters(workbook) -> list[dict[str, Any]]:
    if "角色名" not in workbook.sheetnames:
        return []

    worksheet = workbook["角色名"]
    rows = [[clean(cell.value) for cell in row] for row in worksheet.iter_rows()]
    characters: list[dict[str, Any]] = []
    if len(rows) >= 2:
        header_character = import_header_character(worksheet, rows[0], rows[1])
        if header_character is not None:
            characters.append(header_character)

    for row_number, row in enumerate(rows[2:], start=3):
        character = import_row_character(worksheet, row_number, row)
        if character is not None:
            characters.append(character)
    return characters


def convert_workbook(input_path: Path) -> dict[str, Any]:
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    return {
        "source": {
            "type": "legacy_artist_workbook",
            "path": str(input_path),
        },
        "artistChains": import_artist_chains(workbook),
        "characters": import_characters(workbook),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    data = convert_workbook(args.input)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(data, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(
        f"Wrote {len(data['artistChains'])} artist chains and "
        f"{len(data['characters'])} characters to {args.output}"
    )


if __name__ == "__main__":
    main()
